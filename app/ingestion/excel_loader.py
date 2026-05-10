from pathlib import Path

from openpyxl import load_workbook

from app.extraction.schema import SourceChunk
from app.ingestion.base import BaseLoader


class ExcelLoader(BaseLoader):
    source_type = "xlsx"

    def supports(self, file_path: Path) -> bool:
        return file_path.suffix.lower() in {".xlsx", ".xlsm"}

    def load(self, file_path: Path) -> list[SourceChunk]:
        formula_workbook = load_workbook(file_path, data_only=False, read_only=True)
        value_workbook = load_workbook(file_path, data_only=True, read_only=True)
        chunks: list[SourceChunk] = []

        try:
            for sheet_name in formula_workbook.sheetnames:
                formula_sheet = formula_workbook[sheet_name]
                value_sheet = value_workbook[sheet_name]

                for row_number, formula_row in enumerate(formula_sheet.iter_rows(), start=1):
                    value_row = next(
                        value_sheet.iter_rows(
                            min_row=row_number,
                            max_row=row_number,
                        )
                    )
                    cells = [
                        _cell_to_text(formula_cell.value, value_cell.value, formula_cell.coordinate)
                        for formula_cell, value_cell in zip(formula_row, value_row)
                    ]
                    text = " | ".join(cell for cell in cells if cell)
                    if not text:
                        continue

                    location = f"sheet {sheet_name} row {row_number}"
                    chunks.append(
                        SourceChunk(
                            chunk_id=_chunk_id(file_path, location),
                            source_file=file_path.name,
                            source_type=self.source_type,
                            source_location=location,
                            content_type="worksheet_row",
                            text=text,
                            metadata={
                                "sheet_name": sheet_name,
                                "row_number": row_number,
                                "max_column": formula_sheet.max_column,
                            },
                        )
                    )
        finally:
            formula_workbook.close()
            value_workbook.close()

        return chunks


def _cell_to_text(formula_value, visible_value, coordinate: str) -> str:
    if formula_value is None and visible_value is None:
        return ""

    if isinstance(formula_value, str) and formula_value.startswith("="):
        if visible_value is None:
            return f"{coordinate}: {formula_value}"
        return f"{coordinate}: {visible_value} ({formula_value})"

    value = visible_value if visible_value is not None else formula_value
    return f"{coordinate}: {value}"


def _chunk_id(file_path: Path, location: str) -> str:
    return f"{file_path.name}:{location}"
