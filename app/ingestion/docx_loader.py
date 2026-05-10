from pathlib import Path

from docx import Document

from app.extraction.schema import SourceChunk
from app.ingestion.base import BaseLoader


class DocxLoader(BaseLoader):
    source_type = "docx"

    def supports(self, file_path: Path) -> bool:
        return file_path.suffix.lower() == ".docx"

    def load(self, file_path: Path) -> list[SourceChunk]:
        document = Document(file_path)
        chunks: list[SourceChunk] = []

        for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()
            if not text:
                continue

            location = f"paragraph {paragraph_number}"
            chunks.append(
                SourceChunk(
                    chunk_id=_chunk_id(file_path, location),
                    source_file=file_path.name,
                    source_type=self.source_type,
                    source_location=location,
                    content_type="paragraph",
                    text=text,
                    metadata={
                        "paragraph_number": paragraph_number,
                        "style": paragraph.style.name if paragraph.style else None,
                    },
                )
            )

        for table_number, table in enumerate(document.tables, start=1):
            for row_number, row in enumerate(table.rows, start=1):
                cell_values = [_clean_text(cell.text) for cell in row.cells]
                text = " | ".join(value for value in cell_values if value)
                if not text:
                    continue

                location = f"table {table_number} row {row_number}"
                chunks.append(
                    SourceChunk(
                        chunk_id=_chunk_id(file_path, location),
                        source_file=file_path.name,
                        source_type=self.source_type,
                        source_location=location,
                        content_type="table_row",
                        text=text,
                        metadata={
                            "table_number": table_number,
                            "row_number": row_number,
                            "cell_count": len(row.cells),
                        },
                    )
                )

        return chunks


def _clean_text(text: str) -> str:
    return " ".join(text.split())


def _chunk_id(file_path: Path, location: str) -> str:
    return f"{file_path.name}:{location}"
