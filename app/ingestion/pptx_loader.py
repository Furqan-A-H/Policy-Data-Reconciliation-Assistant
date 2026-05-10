from pathlib import Path

from pptx import Presentation

from app.extraction.schema import SourceChunk
from app.ingestion.base import BaseLoader


class PptxLoader(BaseLoader):
    source_type = "pptx"

    def supports(self, file_path: Path) -> bool:
        return file_path.suffix.lower() == ".pptx"

    def load(self, file_path: Path) -> list[SourceChunk]:
        presentation = Presentation(file_path)
        chunks: list[SourceChunk] = []

        for slide_number, slide in enumerate(presentation.slides, start=1):
            slide_text_parts: list[str] = []
            table_number = 0

            for shape in slide.shapes:
                if getattr(shape, "has_table", False):
                    table_number += 1
                    table_text = _table_to_text(shape.table)
                    if table_text:
                        location = f"slide {slide_number} table {table_number}"
                        chunks.append(
                            SourceChunk(
                                chunk_id=_chunk_id(file_path, location),
                                source_file=file_path.name,
                                source_type=self.source_type,
                                source_location=location,
                                content_type="table",
                                text=table_text,
                                metadata={
                                    "slide_number": slide_number,
                                    "table_number": table_number,
                                    "row_count": len(shape.table.rows),
                                    "column_count": len(shape.table.columns),
                                },
                            )
                        )
                    continue

                if getattr(shape, "has_text_frame", False):
                    shape_text = _clean_text(shape.text)
                    if shape_text:
                        slide_text_parts.append(shape_text)

            slide_text = "\n".join(slide_text_parts).strip()
            if slide_text:
                location = f"slide {slide_number}"
                chunks.append(
                    SourceChunk(
                        chunk_id=_chunk_id(file_path, location),
                        source_file=file_path.name,
                        source_type=self.source_type,
                        source_location=location,
                        content_type="slide_text",
                        text=slide_text,
                        metadata={
                            "slide_number": slide_number,
                            "text_shape_count": len(slide_text_parts),
                        },
                    )
                )

        return chunks


def _table_to_text(table) -> str:
    rows: list[str] = []
    for row in table.rows:
        values = [_clean_text(cell.text) for cell in row.cells]
        row_text = " | ".join(value for value in values if value)
        if row_text:
            rows.append(row_text)
    return "\n".join(rows)


def _clean_text(text: str) -> str:
    return " ".join(text.split())


def _chunk_id(file_path: Path, location: str) -> str:
    return f"{file_path.name}:{location}"
