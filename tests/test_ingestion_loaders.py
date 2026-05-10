import logging

import pytest

from app.ingestion.base import load_files_from_directory

docx = pytest.importorskip("docx")
openpyxl = pytest.importorskip("openpyxl")
pptx = pytest.importorskip("pptx")
from pptx.util import Inches

from app.ingestion.docx_loader import DocxLoader
from app.ingestion.excel_loader import ExcelLoader
from app.ingestion.pptx_loader import PptxLoader


def test_docx_loader_extracts_paragraphs_and_tables(tmp_path) -> None:
    document = docx.Document()
    document.add_paragraph("Affordable housing completions increased.")
    document.add_paragraph("")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Homes delivered"
    table.cell(1, 1).text = "120"

    file_path = tmp_path / "briefing.docx"
    document.save(file_path)

    chunks = DocxLoader().load(file_path)

    assert [chunk.source_location for chunk in chunks] == [
        "paragraph 1",
        "table 1 row 1",
        "table 1 row 2",
    ]
    assert chunks[0].source_type == "docx"
    assert "Affordable housing completions" in chunks[0].text
    assert chunks[2].metadata["row_number"] == 2


def test_pptx_loader_extracts_slide_text_and_tables(tmp_path) -> None:
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Housing supply"
    table_shape = slide.shapes.add_table(
        rows=2,
        cols=2,
        left=Inches(1),
        top=Inches(1.5),
        width=Inches(5),
        height=Inches(1),
    )
    table = table_shape.table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Completions"
    table.cell(1, 1).text = "85"

    file_path = tmp_path / "deck.pptx"
    presentation.save(file_path)

    chunks = PptxLoader().load(file_path)

    locations = {chunk.source_location for chunk in chunks}
    assert "slide 1" in locations
    assert "slide 1 table 1" in locations
    assert any("Housing supply" in chunk.text for chunk in chunks)
    assert any("Completions | 85" in chunk.text for chunk in chunks)


def test_excel_loader_extracts_rows_and_formula_text(tmp_path) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Summary"
    worksheet["A1"] = "Metric"
    worksheet["B1"] = "Value"
    worksheet["A2"] = "Affordable homes"
    worksheet["B2"] = 100
    worksheet["C2"] = "=B2*2"

    file_path = tmp_path / "metrics.xlsx"
    workbook.save(file_path)

    chunks = ExcelLoader().load(file_path)

    assert [chunk.source_location for chunk in chunks] == [
        "sheet Summary row 1",
        "sheet Summary row 2",
    ]
    assert chunks[1].source_type == "xlsx"
    assert "B2: 100" in chunks[1].text
    assert "C2: =B2*2" in chunks[1].text


def test_load_files_from_directory_continues_after_bad_file(tmp_path, caplog) -> None:
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "Housing supply"
    workbook_path = tmp_path / "metrics.xlsx"
    workbook.save(workbook_path)

    bad_docx = tmp_path / "broken.docx"
    bad_docx.write_text("not a real docx")

    with caplog.at_level(logging.ERROR):
        chunks = load_files_from_directory(tmp_path)

    assert len(chunks) == 1
    assert chunks[0].source_file == "metrics.xlsx"
    assert "Failed to load" in caplog.text
